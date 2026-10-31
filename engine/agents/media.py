"""Agent 8 - Media. Picks the showcase from the venture's industry and builds the assets for real.

Never reuses client media. Charts are drawn from the computed series; diagrams are redrawn from a
generic component list. Every label goes through leak_scan, every claim on the slide goes through the
fact-checker, and every image is re-saved with no metadata."""
from __future__ import annotations

import html
import json
import math

from ..core import RunContext, md_table
from .analyst import series
from .exposure import leak_scan
from .factcheck import check_line

PLAYBOOK = {
    "manufacturing": ["cohort_chart", "phase_chart", "architecture", "one_slide", "roi_calculator"],
    "software": ["architecture", "phase_chart", "one_slide"],
    "ai_data_services": ["architecture", "one_slide"],
    "public_sector": ["phase_chart", "one_slide"],
    "healthcare": ["phase_chart", "one_slide"],
}
INK, MUTED, ACCENT, ACCENT2, GRID = "#1f2933", "#6b7280", "#1d6fb8", "#d97706", "#e5e7eb"


def clean_png(path):
    from PIL import Image
    with Image.open(path) as im:
        data = im.convert("RGB")
    data.save(path, format="PNG")  # no info dict, no EXIF, no text chunks


def check_labels(ctx: RunContext, labels: list[str], asset: str) -> None:
    for lab in labels:
        hits = leak_scan(lab, ctx)
        if hits:
            raise ValueError(f"media {asset}: label '{lab}' would expose the client: {hits}")


def style(ax, title, subtitle):
    ax.spines[["top", "right"]].set_visible(False)
    ax.spines[["left", "bottom"]].set_color(MUTED)
    ax.tick_params(colors=MUTED, labelsize=9)
    ax.grid(axis="y", color=GRID, linewidth=0.8)
    ax.set_axisbelow(True)
    ax.figure.text(0.06, 0.95, title, fontsize=14, fontweight="bold", color=INK)
    ax.figure.text(0.06, 0.905, subtitle, fontsize=9.5, color=MUTED)


def cohort_chart(ctx, spec, path):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    fig, ax = plt.subplots(figsize=(10, 5.4), dpi=150)
    fig.subplots_adjust(top=0.82, bottom=0.17, left=0.07, right=0.97)
    weeks_all = sorted({k for s in spec["series"] for k, _ in series(ctx, s)})
    idx = {w: i + 1 for i, w in enumerate(weeks_all)}
    labels = [spec["title"], spec["subtitle"], spec["footer"], spec["x_label"], spec["y_label"]]
    for i, (s, colour) in enumerate(zip(spec["series"], [ACCENT, ACCENT2, INK])):
        pts = series(ctx, s)
        ax.plot([idx[k] for k, _ in pts], [v for _, v in pts], marker="o", ms=4, lw=2.2, color=colour, label=s["label"], zorder=10 - i)
        x, y = idx[pts[-1][0]], pts[-1][1]
        ax.annotate(f"{y:.1f}%", (x, y), xytext=(6, 0), textcoords="offset points", va="center", fontsize=9, color=colour)
        labels.append(s["label"])
    for ann in spec.get("annotations", []):
        ax.axvline(idx[ann["at"]] - 0.5, color=MUTED, ls=":", lw=1)
        ax.text(idx[ann["at"]] - 0.4, spec["ylim"][0] + 1.5, ann["text"], fontsize=8, color=MUTED, rotation=90, va="bottom")
        labels.append(ann["text"])
    if spec.get("target"):
        ax.axhline(spec["target"]["value"], color=INK, lw=0.8, ls="--")
        ax.text(1, spec["target"]["value"] + 0.6, spec["target"]["label"], fontsize=8, color=INK)
        labels.append(spec["target"]["label"])
    check_labels(ctx, labels, "cohort_chart")
    ax.set_ylim(*spec["ylim"])
    ax.set_xlim(0.5, len(weeks_all) + 1.5)
    ax.set_xticks(range(1, len(weeks_all) + 1))
    ax.set_xlabel(spec["x_label"], color=MUTED, fontsize=9)
    ax.set_ylabel(spec["y_label"], color=MUTED, fontsize=9)
    ax.legend(frameon=False, fontsize=9, loc="lower right")
    style(ax, spec["title"], spec["subtitle"])
    fig.text(0.06, 0.03, spec["footer"], fontsize=7.5, color=MUTED)
    fig.savefig(path, metadata={"Software": None})
    plt.close(fig)
    clean_png(path)


def phase_chart(ctx, spec, path):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    pts = series(ctx, spec["series"])
    fig, ax = plt.subplots(figsize=(10, 5.4), dpi=150)
    fig.subplots_adjust(top=0.82, bottom=0.2, left=0.07, right=0.97)
    days = list(range(1, len(pts) + 1))
    vals = [v for _, v in pts]
    labels = [spec["title"], spec["subtitle"], spec["footer"], spec["x_label"], spec["y_label"]]
    keys = [k for k, _ in pts]
    for ph, colour in zip(spec["phases"], ["#9ca3af", ACCENT, "#60a5fa", ACCENT2]):
        a, b = keys.index(ph["from"]), keys.index(ph["to"])
        xs = [d for d, v in zip(days[a:b + 1], vals[a:b + 1]) if not math.isnan(v)]
        ys = [v for v in vals[a:b + 1] if not math.isnan(v)]
        ax.plot(xs, ys, lw=1.6, color=colour, marker="o" if ph.get("dots") else None, ms=2.5)
        mean_v = ctx.metrics[ph["metric"]]["value"]
        ax.hlines(mean_v, xs[0], xs[-1], color=colour, lw=3, alpha=0.9, label=f"{ph['label']}: average {mean_v:.0f} min")
        labels.append(ph["label"])
    ax.legend(frameon=False, fontsize=9, loc="lower left")
    check_labels(ctx, labels, "phase_chart")
    ax.set_ylim(*spec["ylim"])
    ax.set_xlabel(spec["x_label"], color=MUTED, fontsize=9)
    ax.set_ylabel(spec["y_label"], color=MUTED, fontsize=9)
    style(ax, spec["title"], spec["subtitle"])
    fig.text(0.06, 0.03, spec["footer"], fontsize=7.5, color=MUTED, wrap=True)
    fig.savefig(path, metadata={"Software": None})
    plt.close(fig)
    clean_png(path)


def architecture(ctx, spec, path):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.patches import FancyBboxPatch
    fig, ax = plt.subplots(figsize=(11, 6), dpi=150)
    ax.set_xlim(0, 100)
    ax.set_ylim(0, 60)
    ax.axis("off")
    labels = [spec["title"], spec["subtitle"], spec["footer"], spec["boundary"]]
    ax.add_patch(FancyBboxPatch((2, 4), 70, 46, boxstyle="round,pad=0.4,rounding_size=1.5", fc="#f3f7fb", ec=ACCENT, lw=1.4, ls="--"))
    ax.text(4, 47.5, spec["boundary"], fontsize=9, color=ACCENT, fontweight="bold")
    boxes = {}
    for b in spec["boxes"]:
        x, y, w, h = b["xywh"]
        fc = "#ffffff" if b.get("inside", True) else "#fff7ed"
        ec = INK if b.get("inside", True) else ACCENT2
        ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.3,rounding_size=1", fc=fc, ec=ec, lw=1.2))
        ax.text(x + 1.2, y + h - 2.6, b["title"], fontsize=9.5, fontweight="bold", color=INK, va="top")
        ax.text(x + 1.2, y + h - 6.2, b["body"], fontsize=8, color=MUTED, va="top", linespacing=1.4)
        boxes[b["id"]] = (x, y, w, h)
        labels += [b["title"], b["body"]]
    for a, b, lab in spec["arrows"]:
        (x1, y1, w1, h1), (x2, y2, w2, h2) = boxes[a], boxes[b]
        if abs((y1 + h1 / 2) - (y2 + h2 / 2)) < 1:  # same row: right edge -> left edge
            p1, p2 = (x1 + w1 + 0.6, y1 + h1 / 2), (x2 - 0.6, y2 + h2 / 2)
        else:  # different rows: top edge -> bottom edge
            p1, p2 = (x1 + w1 / 2, y1 + h1 + 0.6), (x2 + w2 / 2, y2 - 0.6)
        ax.annotate("", xy=p2, xytext=p1, arrowprops=dict(arrowstyle="-|>", color=MUTED, lw=1.2))
        if lab:
            ax.text((p1[0] + p2[0]) / 2, (p1[1] + p2[1]) / 2 + 1.2, lab, fontsize=7, color=MUTED, ha="center", va="bottom", linespacing=1.1)
            labels.append(lab)
    check_labels(ctx, labels, "architecture")
    fig.text(0.04, 0.95, spec["title"], fontsize=14, fontweight="bold", color=INK)
    fig.text(0.04, 0.915, spec["subtitle"], fontsize=9.5, color=MUTED)
    fig.text(0.04, 0.03, spec["footer"], fontsize=7.5, color=MUTED)
    fig.savefig(path, metadata={"Software": None})
    plt.close(fig)
    clean_png(path)


def one_slide(ctx, spec, path, chart_path):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt
    lines = []
    for raw in [spec["headline"]] + spec["bullets"] + [spec["footer"]]:
        ok, reason, clean = check_line(raw, ctx)
        if not ok:
            raise ValueError(f"slide line refused: {clean} - {reason}")
        lines.append(clean)
    headline, bullets, footer = lines[0], lines[1:-1], lines[-1]
    check_labels(ctx, [spec["kicker"]], "one_slide")
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    s = prs.slides.add_slide(prs.slide_layouts[6])

    def text(x, y, w, h, t, size, bold=False, colour=INK):
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = t
        p.font.size, p.font.bold = Pt(size), bold
        p.font.color.rgb = RGBColor.from_string(colour.lstrip("#"))
        return tf
    text(0.6, 0.35, 12, 0.4, spec["kicker"], 13, colour=ACCENT)
    text(0.6, 0.75, 12.2, 1.2, headline, 30, bold=True)
    tf = text(0.6, 2.2, 5.6, 4.2, "", 15)
    for i, b in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(15)
        p.space_after = Pt(10)
    s.shapes.add_picture(str(chart_path), Inches(6.5), Inches(2.1), width=Inches(6.4))
    text(0.6, 6.75, 12, 0.6, footer, 10, colour=MUTED)
    prs.core_properties.author = ""
    prs.core_properties.title = "Case study slide (draft, pending client approval)"
    prs.save(path)


def roi_calculator(ctx, spec, path):
    fields = spec["fields"]
    check_labels(ctx, [spec["title"], spec["intro"], spec["note"]] + [f["label"] for f in fields], "roi_calculator")
    ok, reason, measured = check_line(spec["measured_line"], ctx)
    if not ok:
        raise ValueError(f"calculator measured line refused: {reason}")
    inputs = "\n".join(
        f'<label>{html.escape(f["label"])}<input id="{f["id"]}" type="number" step="any" value="{f.get("default", "")}" placeholder="{html.escape(f.get("placeholder", ""))}"></label>'
        for f in fields)
    page = f"""<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Gate payback calculator</title><style>
:root{{--bg:#fff;--ink:#1f2933;--muted:#6b7280;--accent:#1d6fb8;--card:#f3f7fb}}
@media (prefers-color-scheme:dark){{:root{{--bg:#111827;--ink:#e5e7eb;--muted:#9ca3af;--accent:#60a5fa;--card:#1f2937}}}}
body{{background:var(--bg);color:var(--ink);font:16px/1.5 system-ui,sans-serif;margin:0;padding:24px 16px}}
main{{max-width:680px;margin:auto}}h1{{font-size:24px;margin:0 0 8px}}p{{color:var(--muted)}}
label{{display:block;margin:14px 0;font-weight:600}}input{{display:block;width:100%;box-sizing:border-box;padding:10px;font:inherit;margin-top:4px;border:1px solid var(--muted);border-radius:6px;background:var(--bg);color:var(--ink)}}
.out{{background:var(--card);border-radius:10px;padding:16px;margin-top:20px}}.big{{font-size:28px;font-weight:700;color:var(--accent)}}small{{color:var(--muted)}}
</style></head><body><main>
<h1>{html.escape(spec["title"])}</h1><p>{html.escape(spec["intro"])}</p>
<p><strong>Measured on the reference project:</strong> {html.escape(measured)}</p>
{inputs}
<div class="out"><div>Estimated monthly saving attributed to the system</div><div class="big" id="save">-</div>
<div>Simple payback</div><div class="big" id="payback">-</div></div>
<p><small>{html.escape(spec["note"])}</small></p>
<script>
const g=id=>parseFloat(document.getElementById(id).value);
function calc(){{const c=g('cost'),r=g('reduction')/100,s=g('share')/100,p=g('project');
const save=c*r*s;document.getElementById('save').textContent=isFinite(save)?save.toLocaleString('en-IN',{{maximumFractionDigits:0}})+' per month':'-';
const pb=p/save;document.getElementById('payback').textContent=isFinite(pb)&&pb>0?pb.toFixed(1)+' months':'-';}}
document.querySelectorAll('input').forEach(i=>i.addEventListener('input',calc));calc();
</script></main></body></html>"""
    path.write_text(page, encoding="utf-8")


def run(ctx: RunContext) -> list:
    cfg = ctx.config
    plan = PLAYBOOK.get(cfg["industry"], PLAYBOOK["manufacturing"])
    media_cfg = cfg["media"]
    out_dir = ctx.out_dir / "05_media"
    out_dir.mkdir(parents=True, exist_ok=True)
    made = []
    builders = {
        "cohort_chart": lambda s: cohort_chart(ctx, s, out_dir / s["file"]),
        "phase_chart": lambda s: phase_chart(ctx, s, out_dir / s["file"]),
        "architecture": lambda s: architecture(ctx, s, out_dir / s["file"]),
        "one_slide": lambda s: one_slide(ctx, s, out_dir / s["file"], out_dir / media_cfg["cohort_chart"]["file"]),
        "roi_calculator": lambda s: roi_calculator(ctx, s, out_dir / s["file"]),
    }
    for kind in plan:
        if kind not in media_cfg:
            continue
        spec = media_cfg[kind]
        builders[kind](spec)
        made.append({"asset": kind, "file": f"05_media/{spec['file']}", "label": spec["label"], "why": spec["why"],
                     "reader": spec["reader"], "sources": spec.get("sources", [])})
    ctx.media = made
    md = ["# Media pack", "", f"Industry: **{cfg['industry']}** - playbook: {', '.join(plan)}.",
          "No client photo, screenshot or drawing is reused; see the exposure report section 4 for why each original was rejected. "
          "All images were re-saved with no metadata. Every label passed the leak scan.", "",
          md_table(["Asset", "File", "Label shown on it", "Reader", "Why this asset", "Built from"],
                   [[m["asset"], f"`{m['file']}`", m["label"], m["reader"], m["why"], ", ".join(m["sources"])] for m in made]), ""]
    ctx.write("05_media/README.md", "\n".join(md))
    ctx.write_json("work/media_manifest.json", made)
    ctx.note("media", f"built {len(made)} assets: {', '.join(m['asset'] for m in made)}")
    return made

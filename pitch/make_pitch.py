"""Builds pitch/pitch.pptx (3 slides). Run after the engine: python pitch/make_pitch.py"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
MEDIA = ROOT / "projects/drishti/outputs/05_media"
INK, MUTED, ACCENT = "1F2933", "6B7280", "1D6FB8"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)


def text(slide, x, y, w, h, lines, size=16, bold=False, colour=INK):
    tf = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
    tf.word_wrap = True
    for i, line in enumerate([lines] if isinstance(lines, str) else lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.size, p.font.bold = Pt(size), bold
        p.font.color.rgb = RGBColor.from_string(colour)
        p.space_after = Pt(8)


# 1. Problem
s = prs.slides.add_slide(prs.slide_layouts[6])
text(s, 0.6, 0.4, 12, 0.4, "THE PROBLEM", 13, colour=ACCENT)
text(s, 0.6, 0.8, 12, 1.4, "Our best proof is locked in email threads - and the easy version of it is wrong", 32, True)
text(s, 0.6, 2.5, 6, 4.5, [
    "The sales draft for this project said: 74% → 97% compliance, zero accidents, ₹2.3 Cr saved.",
    "Every one of those fails against the record: mixed baselines, seven months of no injuries claimed as a result, savings the client's finance team had already rejected.",
    "The client said no to the WhatsApp approval, and warned that 'cement major' with those numbers names them.",
    "Today a person spends days digging, and still gets it wrong."], 17)
text(s, 7.0, 2.5, 5.8, 4.5, [
    "What a buyer, a plant team and a legal team need instead:",
    "• numbers measured the way the client signed for",
    "• missed targets shown as missed, and how they were fixed",
    "• credit shared with what the client did",
    "• nothing that identifies the client, alone or combined"], 17, colour=MUTED)

# 2. Engine
s = prs.slides.add_slide(prs.slide_layouts[6])
text(s, 0.6, 0.4, 12, 0.4, "THE ENGINE", 13, colour=ACCENT)
text(s, 0.6, 0.8, 12, 1.0, "Nine agents, one rule: no line leaves without a file behind it", 30, True)
text(s, 0.6, 2.0, 12.2, 1.2, "Librarian → Extractor → Historian → Analyst → Exposure officer → Writer → Fact-checker → Media → Approval clerk → STOP for a person", 18, colour=INK)
text(s, 0.6, 3.1, 6, 4, [
    "• 157 evidence items, every quote verified word-for-word to file:line",
    "• 54 metrics recomputed from raw data; reproduces the client's acceptance certificate exactly (91.1%, −46%, 4.9 min, 9.7%)",
    "• Refuses: untraceable numbers, ruled-out claims (zero LTIs, 97%), identifiers (names, places, capacity, camera count, weighbridge press release, GPS in a photo)",
    "• Two signed approvals, tied to the pack's hash; plant approval alone does not count"], 16)
text(s, 7.0, 3.1, 5.8, 4, [
    "Runs on another project untouched: only project.json changes (tested on a second, different library).",
    "9 automated checks pass, incl. a refusal test and a stop-for-a-person test.",
    "Claude writes and extracts with --llm; the checker, calculator and gate are code, so the model can propose but never publish."], 16, colour=MUTED)

# 3. Result
s = prs.slides.add_slide(prs.slide_layouts[6])
text(s, 0.6, 0.4, 12, 0.4, "THE RESULT, READY FOR A YES", 13, colour=ACCENT)
text(s, 0.6, 0.8, 7, 1.4, "Real PPE compliance 69% → 91% in an eight-week pilot - and the story a client can sign", 26, True)
text(s, 0.6, 2.6, 6, 4.4, [
    "Gatekeeper: one-page fact sheet - no video leaves the plant, text-only alerts, outage owned and fixed.",
    "User: 'I showed the photo. After that, no argument.' False alarms 20% → under 10%, alerts acknowledged in ~2 min.",
    "Buyer: gate turnaround −24% like-for-like, money only where finance measured it, payback in the client's own words.",
    "Approval: anonymised / named later / internal only / no - reply by 18 Nov, withdraw any time."], 15)
chart = MEDIA / "chart_compliance_cohorts.png"
if chart.exists():
    s.shapes.add_picture(str(chart), Inches(6.9), Inches(2.4), width=Inches(6.1))

prs.save(ROOT / "pitch/pitch.pptx")
print("wrote pitch/pitch.pptx")
